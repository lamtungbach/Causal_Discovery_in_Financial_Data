"""Build the editable and submission versions of the six-page KSE paper."""

from __future__ import annotations

from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import networkx as nx
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / "results" / "camera_ready" / "tables"
OUT = ROOT / "output" / "paper"
ASSETS = OUT / "assets"
DOCX_PATH = OUT / "KSE2026_CameraReady_Revised_6Pages.docx"

FONT = "Times New Roman"
BODY_SIZE = 10.0
CAPTION_SIZE = 8.0
COLUMN_TWIPS = 5020


def set_font(run, size: float, *, bold=False, italic=False, name=FONT) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_margins(cell, top=35, start=40, bottom=35, end=40) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        row._tr.get_or_add_trPr()
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for child in list(borders):
        borders.remove(child)
    for edge in ("top", "bottom", "left", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    for edge in ("insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = tc_pr.first_child_found_in("w:tcBorders")
        if tc_borders is None:
            tc_borders = OxmlElement("w:tcBorders")
            tc_pr.append(tc_borders)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "000000")
        tc_borders.append(bottom)


def configure_section(section, columns: int) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(13)
    section.bottom_margin = Mm(14)
    section.left_margin = Mm(13)
    section.right_margin = Mm(13)
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)
    sect_pr = section._sectPr
    cols = sect_pr.first_child_found_in("w:cols")
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "360")
    cols.set(qn("w:equalWidth"), "1")


def set_keep(paragraph, *, next_=False, together=False) -> None:
    pf = paragraph.paragraph_format
    pf.keep_with_next = next_
    pf.keep_together = together


def add_body(doc: Document, text: str, *, no_indent=False) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0.7)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.first_line_indent = None if no_indent else Inches(0.16)
    set_keep(p, together=False)
    set_font(p.add_run(text), BODY_SIZE)


def add_heading(doc: Document, number: str, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    set_keep(p, next_=True)
    label = f"{number}.  {title}" if number else title
    run = p.add_run(label)
    set_font(run, 10.0)
    run.font.small_caps = True


def add_subheading(doc: Document, letter: str, title: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(1)
    set_keep(p, next_=True)
    set_font(p.add_run(f"{letter}.  {title}"), 9.7, italic=True)


def add_equation(doc: Document, equation: str, number: int) -> None:
    """Add a marker that Word converts to a native professional equation."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.12))
    run = p.add_run(f"[[MATH:{number}]]{equation}[[/MATH]]")
    set_font(run, 9.4, name="Cambria Math")
    set_keep(p, together=True)


def add_caption(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    set_keep(p, next_=True)
    r1 = p.add_run(label.upper())
    set_font(r1, CAPTION_SIZE)
    r1.font.small_caps = True
    p.add_run().add_break()
    r2 = p.add_run(text)
    set_font(r2, CAPTION_SIZE)
    r2.font.small_caps = True


def add_figure_caption(doc: Document, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(f"{label}. "), CAPTION_SIZE)
    set_font(p.add_run(text), CAPTION_SIZE)


def add_table(doc: Document, label: str, caption: str, headers, rows, widths) -> None:
    add_caption(doc, label, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    for j, value in enumerate(headers):
        p = table.rows[0].cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(value)), 7.8)
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(value)), 7.7)
    set_table_geometry(table, widths)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)


def add_image(doc: Document, path: Path, width_inches: float) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(0)
    set_keep(p, next_=True)
    p.add_run().add_picture(str(path), width=Inches(width_inches))


def add_break(doc: Document, break_type) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(break_type)


def generate_figures() -> tuple[Path, Path, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    retained = pd.read_csv(TABLES / "ablation_importance_retained_seed42.csv", index_col=0)
    graph = nx.DiGraph()
    labels = retained.index.tolist()
    graph.add_nodes_from(labels)
    for source in labels:
        for target in labels:
            weight = float(retained.loc[source, target])
            if source != target and weight > 0:
                graph.add_edge(source, target, weight=weight)
    positions = {
        "VNI": (0.00, 0.20),
        "GOLD": (-0.15, -0.68),
        "SSE": (-0.90, 0.35),
        "SP500": (0.88, 0.80),
        "BTC": (0.78, 0.08),
        "VIX": (0.88, -0.70),
        "OIL": (-0.72, -0.80),
    }
    roles = pd.read_csv(TABLES / "node_roles_prediction_weights.csv").set_index("node")
    colors = ["#E69F00" if roles.loc[node, "role"] == "Transmitter" else "#56B4E9" for node in graph]
    fig, ax = plt.subplots(figsize=(3.45, 2.55))
    nx.draw_networkx_nodes(graph, positions, node_size=560, node_color=colors, edgecolors="black", linewidths=0.6, ax=ax)
    nx.draw_networkx_labels(graph, positions, font_size=7.2, font_weight="bold", ax=ax)
    edge_weights = [graph[u][v]["weight"] for u, v in graph.edges]
    nx.draw_networkx_edges(
        graph,
        positions,
        width=[0.8 + 2.2 * w / max(edge_weights) for w in edge_weights],
        edge_color="#333333",
        arrows=True,
        arrowstyle="-|>",
        arrowsize=12,
        connectionstyle="arc3,rad=0.04",
        min_source_margin=13,
        min_target_margin=13,
        ax=ax,
    )
    nx.draw_networkx_edge_labels(
        graph,
        positions,
        edge_labels={(u, v): f"{d['weight']:.3f}" for u, v, d in graph.edges(data=True)},
        font_size=5.7,
        rotate=False,
        label_pos=0.55,
        bbox={"alpha": 0.78, "color": "white", "pad": 0.1},
        ax=ax,
    )
    ax.text(-0.98, -1.02, "Orange: transmitter   Blue: receiver", fontsize=5.8)
    ax.set_xlim(-1.04, 1.06)
    ax.set_ylim(-1.08, 1.02)
    ax.axis("off")
    fig.tight_layout(pad=0.05)
    graph_path = ASSETS / "paper_ngc_graph.png"
    fig.savefig(graph_path, dpi=300, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)

    gbri = pd.read_csv(TABLES / "gbri_prediction_weights.csv", parse_dates=["date"])
    test = gbri.loc[gbri["segment"].eq("test")].dropna(subset=["gbri"])
    fig, ax = plt.subplots(figsize=(3.45, 1.95))
    ax.plot(test["date"], test["gbri"], color="#0072B2", linewidth=0.8)
    for y, color in ((70, "#E69F00"), (85, "#D55E00"), (95, "#9B2226")):
        ax.axhline(y, color=color, linestyle="--", linewidth=0.65)
    ax.set_ylim(0, 101)
    ax.set_ylabel("GBRI percentile", fontsize=7)
    ax.xaxis.set_major_locator(mdates.MonthLocator(interval=4))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%b\n%Y"))
    ax.tick_params(axis="both", labelsize=6.4)
    ax.grid(axis="y", alpha=0.2, linewidth=0.4)
    for spine in ax.spines.values():
        spine.set_linewidth(0.6)
    fig.tight_layout(pad=0.4)
    risk_path = ASSETS / "paper_gbri_test.png"
    fig.savefig(risk_path, dpi=300, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)

    stability = pd.read_csv(TABLES / "edge_stability_20seeds.csv").head(10).copy()
    stability["edge"] = stability["source"] + "->" + stability["target"]
    stability = stability.iloc[::-1]
    fig, ax = plt.subplots(figsize=(3.45, 2.15))
    colors = ["#009E73" if value >= 0.80 else "#999999" for value in stability["selection_frequency"]]
    ax.barh(stability["edge"], stability["selection_frequency"], color=colors, height=0.65)
    ax.axvline(0.80, color="#D55E00", linestyle="--", linewidth=0.8)
    ax.set_xlim(0, 1.02)
    ax.set_xlabel("Selection frequency", fontsize=7)
    ax.tick_params(axis="both", labelsize=6.2)
    ax.grid(axis="x", alpha=0.2, linewidth=0.4)
    for spine in ax.spines.values():
        spine.set_linewidth(0.6)
    fig.tight_layout(pad=0.35)
    stability_path = ASSETS / "paper_edge_stability.png"
    fig.savefig(stability_path, dpi=300, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)
    return graph_path, risk_path, stability_path


def build_document() -> None:
    graph_path, risk_path, stability_path = generate_figures()
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_section(doc.sections[0], 1)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.space_after = Pt(0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(10)
    set_font(
        title.add_run(
            "Identifying Granger-Causal Return Spillovers in Multivariate Financial Data with Interpretable Machine Learning"
        ),
        22.0,
    )
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(2)
    set_font(
        authors.add_run("Nguyen Lam Tung Bach¹, Ngo Minh Huong¹, and Nguyen Van Tang²*"),
        10.5,
        bold=True,
    )
    affiliation = doc.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.paragraph_format.space_after = Pt(5)
    set_font(
        affiliation.add_run(
            "¹Institute for Artificial Intelligence, University of Engineering and Technology,\n"
            "Vietnam National University, Hanoi, Vietnam\n"
            "²Faculty of Technology and Data Science, Foreign Trade University, Hanoi, Vietnam\n"
            "*Corresponding author: tangnv@ftu.edu.vn"
        ),
        9.2,
    )

    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(body_section, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Abstract-")
    set_font(r, 8.6, bold=True, italic=True)
    abstract = (
        "This study examines directed predictive return spillovers among the VN-Index and six global financial series, "
        "and distinguishes them from a volatility-oriented monitoring index. Daily data from January 2015 to May 2026 "
        "are processed under an auditable VN-Index calendar, with explicit treatment of the negative WTI price. A "
        "multivariate VAR(20), an LSTM, and target-specific group-norm-regularized MLPs are evaluated chronologically. "
        "To avoid the scale non-identifiability of raw neural weights, graph edges are scored by validation-loss changes "
        "after complete source-block ablation. Held-out LSTM and Neural Granger forecasts are statistically indistinguishable "
        "under modified Diebold-Mariano and exact McNemar tests. The seed-42 graph retains seven edges, but only VNI->GOLD "
        "meets the pre-specified 0.80 selection-frequency criterion over 20 seeds. VNI is a transmitter in 65% and a receiver "
        "in 35% of runs; its role also changes under common-day sampling. The results therefore support interpretable "
        "predictive links but not a robust claim that Vietnam is uniformly a net receiver. A training-referenced graph-based "
        "risk index is reported as an out-of-sample monitoring score only for the test period."
    )
    set_font(p.add_run(abstract), 8.6, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run("Index Terms-"), 8.6, bold=True, italic=True)
    set_font(
        p.add_run(
            "Granger causality, return spillovers, neural networks, prediction ablation, stability, VN-Index."
        ),
        8.6,
        bold=True,
    )

    add_heading(doc, "I", "Introduction")
    add_body(
        doc,
        "Cross-border equity trading, commodity shocks, digital assets, and global risk sentiment can transmit information to emerging markets. Correlation describes contemporaneous co-movement, whereas Granger causality asks whether one series' lagged history improves prediction of another [1]. Linear vector autoregression (VAR) is transparent [2], but nonlinear and state-dependent relationships motivate neural models."
    )
    add_body(
        doc,
        "Neural Granger Causality (NGC) fits a nonlinear autoregression for each target and groups the lags of each source [3], [4]. Related graph-neural approaches directly model spatial-temporal spillovers for volatility or stock prediction [6], [7]. Our purpose is complementary: the system does not require an ex ante graph, and source exclusion is evaluated from its effect on validation predictions. The learned structure is a model-attribution graph, not structural causality."
    )
    add_body(
        doc,
        "This revision makes four contributions to an applied financial workflow. First, it separates predictive return links from the graph-based risk index (GBRI), which combines shock magnitude with VNI volatility. Second, it audits non-synchronous calendars and market timing. Third, it replaces raw first-layer norms, which are not scale identified in ReLU networks, with prediction-based source-block ablation. Fourth, it reports graph stability, forecast inference, and a multivariate VAR-FDR network baseline. Methods not executed by the verified pipeline, including the earlier multiscale and event-window analyses, are outside the present empirical scope."
    )
    add_body(
        doc,
        "This scope also clarifies terminology. The VAR and neural models predict returns, so their directed edges are predictive return spillovers. Risk transmission is a broader interpretation that requires volatility, loss, or another risk measure [5]. We reserve the risk label for GBRI, where absolute external shocks are combined with domestic volatility. This distinction prevents forecast attribution from being presented as a direct measure of economic loss."
    )
    add_body(
        doc,
        "The empirical questions are therefore: (1) do NGC and LSTM differ in held-out VNI forecast loss or sign accuracy; (2) which source-target relations survive prediction ablation and multiple seeds; (3) how does the NGC network compare with linear VAR-Granger after multiplicity correction; and (4) do VNI's incoming links and node role survive calendar and timing alternatives?"
    )
    add_heading(doc, "II", "Methodology")
    add_subheading(doc, "A", "Data transformation and timing")
    add_body(
        doc,
        "For asset i on date t, the percentage log return is defined only when both endpoint prices are finite and strictly positive:"
    )
    add_equation(doc, "r_(t,i)=100\\ln(P_(t,i)/P_(t-1,i))", 1)
    add_body(
        doc,
        "The main panel is anchored to VNI dates. OIL, GOLD, SP500, and VIX are forward-filled for at most three observations; SSE is forward-filled for at most three and linearly interpolated for at most five; BTC is not filled. WTI's -37.63 close on 20 April 2020 is retained in the price audit, while returns ending on 20 and 21 April are invalidated, preventing a multi-day bridge across the negative price."
    )
    add_body(
        doc,
        "Let x(t,i) standardize r(t,i) using training-only mean and standard deviation. With p=20, X(t)=[x(t-p),...,x(t-1)] is the input and x(t,i) is the target. The forecast origin is immediately before VNI opens on date t. Because a U.S. close labeled t-1 is available before this origin, a sensitivity run shifts OIL, GOLD, SP500, and VIX by one additional VNI observation. A second run uses only dates observed by all seven markets and applies no fill or interpolation."
    )

    add_subheading(doc, "B", "Forecast models and prediction-based edges")
    add_body(
        doc,
        "The LSTM benchmark has one recurrent layer with 64 hidden units and a linear VNI output [9]. NGC fits seven target-specific MLPs with architecture 140-64-64-1 and ReLU activations. For target i, Adam minimizes"
    )
    add_equation(doc, "L_i=MSE_(tr,i)+λ∑_(j=1)^N‖W_(i,j)^((1))‖_F", 2)
    add_body(
        doc,
        "where lambda=5x10^-4 and W(i,j)^(1) contains the 20 inputs from source j. This differentiable group-norm penalty regularizes groups but is not a proximal operator and does not imply exact-zero selection. Validation loss selects each checkpoint; test data are not used for training, early stopping, or graph construction."
    )
    add_body(
        doc,
        "For each off-diagonal source-target pair, all 20 standardized lags of source j are set to zero on validation inputs. With e=10^-12, the nonnegative prediction importance is"
    )
    add_equation(doc, "C_(j→i)=max{0,(MSE_(val)^(-j,i)-MSE_(val)^(i))/(MSE_(val)^(i)+ε)}", 3)
    add_body(
        doc,
        "Signed changes are archived, while positive scores at or above their empirical 0.70 quantile are retained. Thus an edge means that removing the source worsened validation prediction enough to pass a pre-specified sparsity rule; it is neither a p-value nor an economic elasticity."
    )

    add_subheading(doc, "C", "Linear network, inference, stability, and risk index")
    add_body(
        doc,
        "A single multivariate VAR(20) is fitted on the development period, with intercept c, coefficient matrices A(k), and innovation vector u(t):"
    )
    add_equation(doc, "r_t=c+∑_(k=1)^20 A_k r_(t-k)+u_t", 4)
    add_body(
        doc,
        "For each source j and target i, a joint F test evaluates H0: A(k)[i,j]=0 for every k=1,...,20. Benjamini-Hochberg correction is applied across the 42 directed tests at FDR 0.05 [8]. Binary overlap with the NGC graph provides a network-level baseline; p-values are not compared with NGC importance magnitudes."
    )
    add_body(
        doc,
        "Out-of-sample forecasts are compared using the two-sided Harvey-Leybourne-Newbold modified Diebold-Mariano test with squared error [11], [12], exact McNemar testing of paired sign correctness [13], and the Pesaran-Timmermann (PT) directional-predictability test [14]. Graph uncertainty is measured across seeds 42-61. An edge is called stable when selected in at least 80% of runs [16]."
    )
    add_body(
        doc,
        "Across retained edges, outgoing strength is s-out(i)=sum(j!=i) C(i->j), incoming strength is s-in(i)=sum(j!=i) C(j->i), and net strength is their difference. Positive net strength defines a transmitter and negative net strength a receiver. Edge-selection frequency over B=20 seeds is"
    )
    add_equation(doc, "π_(j→i)=1/B ∑_(b=1)^B 1{C̃_(j→i)^((b))>0}", 5)
    add_body(
        doc,
        "External VNI pressure is P(t)=sum(j!=VNI) C(j->VNI)|x(t,j)|. It is combined equally with 20-day VNI return volatility after training-only z-standardization:"
    )
    add_equation(doc, "GBRI_t=100F̂_(tr)[0.5z̃_(P,t)+0.5z̃_(V,t)]", 6)
    add_body(
        doc,
        "Here Ftr is the empirical CDF of the combined statistic in training. Regimes are Low (<70), Medium [70,85), High [85,95), and Extreme (>=95). Training-referenced standardization and percentiles prevent later observations from setting earlier thresholds. Only the test segment is described as out-of-sample monitoring."
    )

    add_subheading(doc, "D", "Reproducible evaluation protocol")
    add_body(
        doc,
        "The pipeline separates three objects that are often conflated: parameter estimation, model selection, and final evaluation. Training observations estimate the return means and standard deviations, fit network parameters, and form the reference distribution used by GBRI. Validation observations select the early-stopping checkpoint and quantify source-block ablation. Test observations are reserved for the forecast scores, paired inference, and prospective GBRI regime counts. Consequently, the graph is an exploratory validation-set object, whereas the reported VNI forecast metrics are held out. This separation also means that training and validation GBRI values are reconstructions and must not be described as contemporaneously available alerts."
    )
    add_body(
        doc,
        "Every sequence is assigned to a split by its target date rather than by the first date in its lag window. The first validation and test sequences may use the preceding 20 observations as historical context, but their targets never enter an earlier split. Standardization parameters are frozen after training. For NGC, seven checkpoints are selected independently because each target has a distinct loss surface; the best epochs range from 2 to 22 and training terminates after 32 to 52 epochs. The seed-42 VNI checkpoint is selected at epoch 14 with standardized validation MSE 0.8002. These details are archived with the predictions so that graph construction can be repeated without retraining."
    )
    add_body(
        doc,
        "Reproducibility is tested at the prediction level, not inferred from identical hyperparameters. On the verified CUDA environment, rerunning seed 42 reproduces all 425 standardized VNI predictions exactly, yielding max_abs_repeat_diff=0.0. This zero is the expected outcome of deterministic execution rather than an error: the prediction standard deviation is 0.1349 and changing to seed 43 changes at least one prediction by 0.7526. We nevertheless treat cross-seed graph variation as epistemic model-selection uncertainty and report it separately from same-seed numerical reproducibility."
    )
    add_body(
        doc,
        "The reported hyperparameters are fixed experiment settings rather than values selected after observing the test set. The lag order p=20 supplies approximately one trading month of history to both neural approaches and the VAR baseline. The NGC penalty coefficient is 5x10^-4, the Adam learning rate is 10^-3, and early stopping uses patience 30 with a maximum of 300 epochs. The seed sweep changes initialization while leaving the data split, preprocessing, architecture, penalty, optimizer, and ablation quantile unchanged. This controlled design attributes graph variability to optimization and finite-sample model selection, although it does not quantify uncertainty over alternative lag orders or regularization levels."
    )

    add_heading(doc, "III", "Data and Experimental Setup")
    add_table(
        doc,
        "Table I",
        "Financial variables and archived data sources",
        ["Symbol", "Market (ticker)", "Provider"],
        [
            ["VNI", "VN-Index", "VCI/vnstock"],
            ["BTC", "Bitcoin (BTC-USD)", "Yahoo Finance"],
            ["OIL", "WTI (CL=F)", "Yahoo Finance"],
            ["GOLD", "Gold (GC=F)", "Yahoo Finance"],
            ["SP500", "S&P 500 (^GSPC)", "Yahoo Finance"],
            ["SSE", "Shanghai (000001.SS)", "Yahoo Finance"],
            ["VIX", "CBOE VIX (^VIX)", "Yahoo Finance"],
        ],
        [700, 2450, 1870],
    )
    add_body(
        doc,
        "The archived files cover January 2015 through 11 May 2026 and have filesystem timestamps of 15 May 2026; the downloader did not embed an API retrieval timestamp. Yahoo Finance files use adjusted close output from yfinance (auto_adjust=True). The main complete price panel has 2,831 rows and produces 2,828 complete return observations from 6 January 2015. The common-day panel has 2,615 returns."
    )
    add_body(
        doc,
        "Splitting is by target date before sequence construction: 1,959 training forecasts (3 February 2015-6 December 2022), 424 validation forecasts (7 December 2022-20 August 2024), and 425 test forecasts (21 August 2024-11 May 2026). Validation and test windows retain the prior 20 observations as context. All seven returns reject the ADF unit-root null at p<0.001; VIX, BTC, and OIL have the largest standard deviations (8.066, 4.238, and 2.982 percentage points)."
    )
    add_body(
        doc,
        "Before calendar completion, the VNI anchor contains 2,847 dates. Missing close counts are 16 for BTC, 103 for OIL, SP500, and VIX, 104 for GOLD, and 140 for SSE. The bounded fill rules recover all but 16 rows, leaving 2,831 complete prices. Return construction then loses the initial difference and the two WTI endpoints, giving 2,828 usable rows. By contrast, the common-day design starts from 2,618 fully observed prices and yields 2,615 returns after applying the same non-positive-price rule. Reporting both designs makes the amount of interpolation and its effect on the inferred role of VNI directly auditable."
    )
    add_table(
        doc,
        "Table II",
        "Descriptive statistics for daily percentage returns",
        ["Series", "Mean", "Std.", "Min", "Max", "ADF"],
        [
            ["VNI", "0.045", "1.169", "-6.910", "6.547", "-49.35"],
            ["BTC", "0.203", "4.238", "-46.473", "22.512", "-10.80"],
            ["OIL", "0.045", "2.982", "-28.221", "33.316", "-9.67"],
            ["GOLD", "0.049", "1.030", "-12.066", "5.905", "-39.66"],
            ["SP500", "0.048", "1.122", "-12.765", "9.089", "-13.04"],
            ["SSE", "0.008", "1.252", "-10.830", "7.755", "-11.31"],
            ["VIX", "-0.009", "8.066", "-44.245", "76.825", "-19.97"],
        ],
        [760, 790, 790, 920, 900, 860],
    )
    add_body(
        doc,
        "The negative OIL minimum in the raw price archive does not appear in Table II because Eq. (1) excludes its two affected return endpoints. The retained return distribution is nevertheless heavy-tailed [17]. The verified implementation uses Python 3.12.3, PyTorch 2.5.1, statsmodels 0.14.6, lambda=5x10^-4, learning rate 10^-3, and seed 42 for the primary fit. Early stopping patience is 30 over at most 300 epochs. Deterministic repetition of seed 42 gives a maximum prediction difference of 0.0 on the verified CUDA stack; seed 43 differs by 0.753, ruling out a constant-output explanation."
    )

    add_heading(doc, "IV", "Results")
    add_subheading(doc, "A", "Held-out forecasting and linear-network baseline")
    add_table(
        doc,
        "Table III",
        "VNI test forecasting performance (n=425)",
        ["Model", "RMSE", "MAE", "R2", "DA (%)"],
        [
            ["LSTM", "1.237", "0.824", "0.0240", "56.00"],
            ["NGC", "1.246", "0.823", "0.0104", "54.59"],
        ],
        [1200, 950, 950, 950, 970],
    )
    add_body(
        doc,
        "LSTM has slightly lower RMSE and higher R2, whereas NGC has slightly lower MAE. The mean NGC-minus-LSTM squared-loss differential is 0.0213; modified DM=0.476 (p=0.635). McNemar discordances are 63 LSTM-only correct signs versus 57 NGC-only correct signs (exact p=0.648). PT statistics are 1.056 for NGC (p=0.291) and 1.645 for LSTM (p=0.100). Hence neither relative forecast superiority nor directional predictability at 5% is supported. NGC's value in this experiment is attribution, not forecast dominance."
    )
    add_table(
        doc,
        "Table IV",
        "Paired forecast and directional inference",
        ["Test", "Statistic/counts", "p-value", "Conclusion"],
        [
            ["Modified DM", "0.476", "0.635", "No loss diff."],
            ["McNemar", "63 vs. 57", "0.648", "No sign diff."],
            ["PT: NGC", "1.056", "0.291", "No direction"],
            ["PT: LSTM", "1.645", "0.100", "No direction"],
        ],
        [1350, 1300, 950, 1420],
    )
    add_body(
        doc,
        "Target-specific NGC forecasts provide an additional scale check. On the test set, RMSE is 2.917 for BTC, 2.784 for OIL, 1.527 for GOLD, 1.049 for SP500, 1.066 for SSE, and 8.770 for VIX. Their R2 values range from -0.039 for SSE to 0.006 for OIL, while directional accuracy ranges from 48.0% to 52.2%. Thus the network is not producing uniformly strong forecasts for the six non-VNI targets. This matters because their fitted models are still used to describe the full directed graph: an edge is evidence of relative prediction dependence within a target model, not proof that the target is accurately predictable in absolute terms."
    )
    add_body(
        doc,
        "The VAR uses 2,383 effective development observations after 20 lags. Eleven of 42 links survive FDR. Incoming VNI links are OIL->VNI (adjusted p=0.0189) and SP500->VNI (adjusted p=5.02x10^-6). The seed-42 NGC graph has seven links and only VIX->VNI. The graphs share one link, VNI->GOLD, among 17 unique links (Jaccard=0.059). Linear significance and nonlinear prediction ablation therefore describe materially different structures."
    )
    add_table(
        doc,
        "Table V",
        "Directed VAR(20) links surviving FDR 0.05",
        ["Source", "Target", "F", "Raw p", "FDR p"],
        [
            ["VNI", "OIL", "2.943", "1.10e-5", "9.24e-5"],
            ["VNI", "GOLD", "2.650", "8.30e-5", "4.98e-4"],
            ["OIL", "VNI", "2.003", "0.00495", "0.01890"],
            ["OIL", "GOLD", "2.277", "9.42e-4", "0.00440"],
            ["OIL", "SP500", "2.794", "3.10e-5", "2.17e-4"],
            ["SP500", "VNI", "3.415", "3.59e-7", "5.02e-6"],
            ["SP500", "OIL", "8.100", "4.28e-24", "1.80e-22"],
            ["SP500", "GOLD", "2.162", "0.00192", "0.00807"],
            ["VIX", "OIL", "3.503", "1.86e-7", "3.91e-6"],
            ["VIX", "SP500", "3.028", "6.02e-6", "6.32e-5"],
            ["VIX", "SSE", "2.592", "1.23e-4", "6.45e-4"],
        ],
        [920, 920, 880, 1150, 1150],
    )
    add_body(
        doc,
        "The full 42-test table is retained in the reproducibility output. Table V lists only FDR discoveries to conserve space. SP500->OIL is the strongest linear result, while VIX has three outgoing FDR links. These patterns provide context for transmitter roles but cannot be read as variance shares because the baseline is coefficient exclusion, not forecast-error variance decomposition."
    )

    add_subheading(doc, "B", "NGC graph, multi-seed stability, and timing sensitivity")
    add_image(doc, graph_path, 3.36)
    add_figure_caption(
        doc,
        "Fig. 1",
        "Seed-42 validation-ablation graph. Weights are relative validation-loss increases, not structural effects.",
    )
    add_table(
        doc,
        "Table VI",
        "Seed-42 retained edges and 20-seed selection frequency",
        ["Edge", "Importance", "Frequency"],
        [
            ["VNI->GOLD", "0.02225", "0.90*"],
            ["BTC->SSE", "0.01660", "0.55"],
            ["SP500->SSE", "0.01228", "0.40"],
            ["GOLD->BTC", "0.00758", "0.10"],
            ["GOLD->SSE", "0.00729", "0.10"],
            ["VIX->VNI", "0.00629", "0.75"],
            ["OIL->SSE", "0.00559", "0.35"],
        ],
        [2050, 1480, 1490],
    )
    add_body(
        doc,
        "The asterisk marks the only edge reaching the 0.80 stability criterion. Two candidate incoming VNI links are nevertheless recurrent: VIX->VNI appears in 75% and SP500->VNI in 65% of runs, although the latter is not retained for seed 42. Edge instability is therefore substantive rather than a numerical-reproducibility failure."
    )
    add_body(
        doc,
        "Across the 20 fits, the empirical 0.70 cutoff ranges from 0.00339 to 0.00790 and each graph retains between five and eight directed edges (mean 6.9). The mean Jaccard similarity with the seed-42 graph is 0.317, including the reference fit itself, and the median is 0.273. Fixed quantile sparsity therefore controls graph size but does not guarantee stable membership. Selection frequency is the more informative summary: VNI->GOLD appears in 18 of 20 runs, whereas no other edge crosses the stated 0.80 criterion."
    )
    add_image(doc, stability_path, 3.30)
    add_figure_caption(
        doc,
        "Fig. 2",
        "Ten most frequently selected NGC edges across seeds 42-61. The dashed line is the pre-specified 0.80 stability criterion.",
    )
    add_table(
        doc,
        "Table VII",
        "Prediction-weighted node roles for the seed-42 graph",
        ["Node", "In", "Out", "Net", "Role"],
        [
            ["VNI", "0.0063", "0.0223", "+0.0160", "Tx"],
            ["SP500", "0.0000", "0.0123", "+0.0123", "Tx"],
            ["BTC", "0.0076", "0.0166", "+0.0090", "Tx"],
            ["VIX", "0.0000", "0.0063", "+0.0063", "Tx"],
            ["OIL", "0.0000", "0.0056", "+0.0056", "Tx"],
            ["GOLD", "0.0223", "0.0149", "-0.0074", "Rx"],
            ["SSE", "0.0418", "0.0000", "-0.0418", "Rx"],
        ],
        [880, 940, 940, 1100, 1160],
    )
    add_body(
        doc,
        "At seed 42, VNI's outgoing VNI->GOLD importance exceeds its sole incoming VIX->VNI importance, so VNI is a transmitter. SSE aggregates four retained inputs and has the most negative net strength. These magnitudes are changes in validation error and should only be compared within the same fitted attribution design."
    )
    add_table(
        doc,
        "Table VIII",
        "Node-role stability across 20 deterministic seeds",
        ["Node", "Tx freq.", "Rx freq.", "Mean net"],
        [
            ["SP500", "0.90", "0.10", "+0.0112"],
            ["VIX", "0.85", "0.10", "+0.0121"],
            ["OIL", "0.80", "0.15", "+0.0076"],
            ["VNI", "0.65", "0.35", "+0.0038"],
            ["BTC", "0.60", "0.15", "+0.0038"],
            ["SSE", "0.15", "0.85", "-0.0167"],
            ["GOLD", "0.05", "0.95", "-0.0218"],
        ],
        [1050, 1250, 1250, 1470],
    )
    add_body(
        doc,
        "SP500, VIX, and OIL are transmitters in at least 80% of runs; GOLD and SSE are predominantly receivers. VNI is a transmitter in 65% and receiver in 35%, contradicting the earlier single-fit receiver claim. This role uncertainty follows from changing retained prediction effects, not from raw neural-weight rescaling."
    )
    add_table(
        doc,
        "Table IX",
        "VNI sensitivity to calendar and global-input timing",
        ["Variant", "n", "Incoming", "Net", "Role"],
        [
            ["Main", "2828", "VIX", "+0.0160", "Tx"],
            ["Common-day", "2615", "VIX", "-0.0018", "Rx"],
            ["Extra shift", "2828", "None", "+0.0149", "Tx"],
        ],
        [1550, 720, 1050, 850, 850],
    )
    add_body(
        doc,
        "The common-day sample removes interpolation and changes VNI to a weak receiver, although VIX remains its incoming link. Shifting global inputs one additional observation removes all incoming VNI edges. Thus lagged global-market associations are sensitive to synchronization and should be described as predictive information flow under a stated timing convention, not invariant economic transmission."
    )
    add_body(
        doc,
        "The magnitude of the role change is also informative. Main-calendar VNI has net strength +0.01596 because outgoing VNI->GOLD exceeds incoming VIX->VNI. In the common-day graph its net strength is -0.00177, a small receiver value rather than a large reversal, while the extra-shift design returns +0.01491 with no retained incoming edge. The direction of the label therefore depends on a net difference close to zero in the common-day case. Reporting the signed strength together with the categorical role avoids presenting a threshold-induced Tx/Rx switch as a discontinuous economic change."
    )

    add_subheading(doc, "C", "Out-of-sample graph-based risk monitoring")
    add_image(doc, risk_path, 3.36)
    add_figure_caption(
        doc,
        "Fig. 3",
        "GBRI over the held-out test period. Dashed lines mark 70, 85, and 95 training-reference percentiles.",
    )
    add_table(
        doc,
        "Table X",
        "Test-period GBRI regime distribution",
        ["Regime", "Range", "Days", "Share (%)"],
        [
            ["Low", "<70", "277", "65.18"],
            ["Medium", "70-85", "75", "17.65"],
            ["High", "85-95", "34", "8.00"],
            ["Extreme", ">=95", "39", "9.18"],
        ],
        [1250, 1250, 1100, 1420],
    )
    add_body(
        doc,
        "The seed-42 main graph gives VIX as the only retained external input to VNI pressure. During the 425-day test period, 17.18% of days are High or Extreme, including 9.18% Extreme compared with a nominal 5% training tail. GBRI is not a loss probability or return-direction forecast: absolute source shocks make it a monitoring intensity. Training and validation values are retrospective reconstructions because the checkpoint and graph were selected with those periods."
    )
    add_body(
        doc,
        "For the training reference, external pressure has mean 0.00435 and standard deviation 0.00455, while 20-day VNI volatility has mean 1.028 and standard deviation 0.523. Both components therefore contribute after separate z-standardization even though their raw units and dispersion differ. The empirical CDF maps the combined score to a percentile without assuming Gaussian tails. Because only VIX enters VNI in the seed-42 graph, the pressure component is sparse in this run; a future operational system should recompute it through rolling-origin estimation rather than permanently fixing one checkpoint and one adjacency matrix."
    )

    add_subheading(doc, "D", "Cross-target forecast audit")
    add_table(
        doc,
        "Table XI",
        "NGC test performance across all target returns (n=425)",
        ["Target", "RMSE", "MAE", "R2", "DA (%)"],
        [
            ["VNI", "1.246", "0.823", "0.010", "54.59"],
            ["BTC", "2.917", "2.125", "-0.013", "48.24"],
            ["OIL", "2.784", "1.872", "0.006", "48.94"],
            ["GOLD", "1.527", "1.036", "-0.036", "52.24"],
            ["SP500", "1.049", "0.681", "-0.001", "52.24"],
            ["SSE", "1.066", "0.681", "-0.039", "48.00"],
            ["VIX", "8.770", "5.894", "-0.011", "48.71"],
        ],
        [1040, 995, 995, 995, 995],
    )
    add_body(
        doc,
        "Table XI reports the target-specific models used to form the directed network, not seven competing forecasts of VNI. RMSE and MAE are in percentage-return units and therefore should not be compared mechanically across assets with different volatility. R2 is evaluated against the test-set mean benchmark. Only VNI and OIL have slightly positive R2, and no target exceeds 54.59% directional accuracy. The graph can still reveal which lag block a fitted target model uses, but these modest predictive scores constrain the strength of any substantive transmission claim."
    )
    add_body(
        doc,
        "This audit also guards against selecting an attractive graph while ignoring weak target forecasts. In particular, SSE receives four seed-42 inputs but has R2=-0.039, so its incoming edges should be read as within-model dependencies rather than evidence of a high-accuracy SSE forecasting system. Similarly, VIX has the largest RMSE because its return scale is much wider than those of the equity indices. Cross-target edge importance is normalized by each target's validation MSE, but normalization does not eliminate differences in data quality, noise, or checkpoint uncertainty."
    )

    add_heading(doc, "V", "Discussion")
    add_subheading(doc, "A", "Cross-method evidence and interpretation")
    add_body(
        doc,
        "The revised evidence changes the substantive conclusion. There is support for lagged predictive information from global variables to VNI, but the identity of the source depends on method and timing: VAR-FDR selects OIL and SP500, seed-42 NGC selects VIX, and the conservative shift selects none. More broadly, SP500, VIX, and OIL frequently act as transmitters, while GOLD and SSE frequently receive prediction-relevant inputs. VNI itself has no stable receiver classification."
    )
    add_body(
        doc,
        "Prediction ablation resolves the specific interpretability problem of first-layer block norms: the score is invariant to compensating rescaling of adjacent ReLU layers because it is calculated from outputs. It remains model- and validation-sample-dependent, however. Zeroing a standardized source is an intervention on the fitted predictor, not on the financial system. Quantile thresholding fixes sparsity rather than statistical significance, so selection frequency is essential."
    )
    add_body(
        doc,
        "Compared with graph-neural spillover models [6], [7], target-specific NGC offers direct source exclusion and does not require an externally specified adjacency matrix. Conversely, graph neural networks may better exploit cross-node parameter sharing and dynamic topology. The present contribution is therefore the audited combination of nonlinear attribution, calendar sensitivity, statistical inference, and a training-referenced monitoring index, rather than a new neural architecture."
    )
    add_body(
        doc,
        "The lack of VAR-NGC edge overlap is not by itself evidence that one method fails. VAR tests joint linear coefficient restrictions on a development sample and controls FDR; NGC measures a nonlinear prediction response on validation data and then enforces fixed sparsity. Their different estimands explain why the graphs need not coincide. Agreement on VNI->GOLD is therefore notable, but one shared edge is too little to claim structural convergence."
    )
    add_body(
        doc,
        "The comparison also separates statistical discovery from predictive relevance. A small VAR p-value concerns a joint coefficient restriction under the linear specification and sampling assumptions; it does not measure effect size. Conversely, a large NGC ablation score means that a fitted nonlinear predictor relies on a source block on the validation sample; it does not supply a null-hypothesis error rate. The FDR-adjusted VAR adjacency and the frequency-filtered NGC adjacency should therefore be viewed as complementary sensitivity analyses rather than interchangeable estimators of a single true network."
    )
    add_body(
        doc,
        "Likewise, the graph and GBRI answer different questions. A directed graph edge says that a lagged source improved a target model under the fitted design. GBRI discards shock direction through absolute source returns and adds domestic volatility, so a high score denotes elevated monitoring conditions. It does not imply that VNI will fall, and the present study does not test a trading strategy."
    )

    add_subheading(doc, "B", "Validity, reproducibility, and deployment")
    add_body(
        doc,
        "The validation set has two explicit roles: it selects the early-stopping checkpoint and supplies the source-ablation loss used to construct the exploratory graph. For that reason, validation graph weights are not presented as held-out performance. Test observations are opened only after model and graph selection for forecast metrics and prospective GBRI labeling. Saved checkpoints reproduce all 425 test predictions, all 42 signed ablation changes are archived, and deterministic same-seed repetition is exact on the reported hardware."
    )
    add_body(
        doc,
        "Limitations include use of daily closes across time zones, interpolation in the main calendar, one fixed lag and lambda, validation reuse for checkpoint and edge selection, and only 20 random seeds without a dependent block bootstrap. The VNI system omits exchange rates, interest rates, bond yields, and capital flows. GBRI also inherits uncertainty from a graph with only one incoming VNI edge. Future work should use rolling-origin graph estimation, nested tuning, block-bootstrap stability, alternative ablation baselines, and formal evaluation against future volatility or drawdown exceedances."
    )
    add_body(
        doc,
        "Three validity boundaries follow. First, temporal precedence at daily resolution is approximate because market closes are asynchronous; the extra-shift result shows that this choice can remove incoming VNI edges. Second, source-block zeroing is an out-of-distribution perturbation only to the extent that a standardized zero represents the training mean; alternative conditional or permutation baselines may change attribution. Third, the test window supplies only 425 forecasts and includes a particular market regime. The paired tests appropriately use common dates, but a non-rejection should be read as insufficient evidence of a difference, not proof that the models are equivalent."
    )

    add_subheading(doc, "C", "Practical interpretation and reporting protocol")
    add_body(
        doc,
        "For research use, each reported link should be accompanied by five qualifiers: the target variable is return; the source timing convention is explicit; importance is computed on validation data; the 0.70 quantile is a sparsity threshold rather than a significance level; and selection frequency summarizes seed sensitivity. Under this protocol, VIX->VNI is a seed-42 exploratory link with frequency 0.75, while OIL->VNI and SP500->VNI are linear FDR discoveries. None of these statements alone establishes a stable structural channel from a global market to Vietnam."
    )
    add_body(
        doc,
        "For monitoring use, GBRI should be treated as a ranked alert that prioritizes dates for analyst review. Its percentile scale is convenient for thresholds, but the observed 9.18% Extreme share does not mean a 9.18% crash probability. A practical dashboard should display the current percentile together with its pressure and domestic-volatility components, the active incoming sources, and the date on which the graph was last estimated. Alerts should be backtested against prespecified future-volatility or drawdown outcomes before they inform capital allocation, hedging, or regulatory action."
    )
    add_body(
        doc,
        "For model maintenance, three diagnostics are especially important. First, graph turnover should be tracked across rolling origins and compared with same-seed numerical checks. Second, calendar coverage and the number of filled prices should be monitored because synchronization changes the inferred VNI role. Third, prediction and attribution drift should be separated: forecast loss can deteriorate even when graph membership appears stable, and graph membership can change while forecast loss remains similar. These diagnostics turn the present static experiment into a falsifiable deployment plan without overstating the evidence available in the current sample."
    )
    add_body(
        doc,
        "The study's strongest contribution is consequently procedural. Each claim is tied to a generated artifact: archived price and missing-value audits support the sample definition; split summaries support the chronology; checkpoint files support model selection; prediction files support the forecast tests; signed ablation matrices support the NGC graph; and seed summaries support stability frequencies. This chain makes disagreements traceable to an estimand, data alignment, or threshold choice. It also prevents a visually persuasive network from substituting for forecast evidence, statistical inference, or robustness analysis."
    )

    add_heading(doc, "VI", "Conclusion")
    add_body(
        doc,
        "This paper presents a reproducible framework for predictive return-spillover analysis and graph-based risk monitoring. It documents data provenance, non-positive WTI handling, target-date splitting, validation checkpoint selection, source-block prediction ablation, multivariate VAR-FDR comparison, forecast inference, 20-seed graph stability, and calendar/timing robustness. NGC does not significantly outperform LSTM, and its seed-42 graph has limited edge stability. The defensible conclusion is not that VNI is uniformly a receiver, but that detected incoming predictive links and VNI's graph role are sensitive to model and synchronization assumptions. GBRI remains a surveillance prototype whose test-period values are genuinely out of sample, not an automated trading rule or evidence of structural causality."
    )

    add_heading(doc, "", "Acknowledgment")
    add_body(
        doc,
        "This work was supported by Vietnam National University, Hanoi, under research project QG.26.68.",
        no_indent=True,
    )

    add_heading(doc, "", "References")
    references = [
        "C. W. J. Granger, 'Investigating causal relations by econometric models and cross-spectral methods,' Econometrica, vol. 37, no. 3, pp. 424-438, 1969.",
        "H. Lutkepohl, New Introduction to Multiple Time Series Analysis. Berlin, Germany: Springer, 2005.",
        "A. Tank, I. Covert, N. Foti, A. Shojaie, and E. B. Fox, 'Neural Granger causality,' IEEE Trans. Pattern Anal. Mach. Intell., vol. 44, no. 8, pp. 4267-4279, 2022.",
        "M. Yuan and Y. Lin, 'Model selection and estimation in regression with grouped variables,' J. Roy. Stat. Soc. B, vol. 68, no. 1, pp. 49-67, 2006.",
        "F. X. Diebold and K. Yilmaz, 'Better to give than to receive: Predictive directional measurement of volatility spillovers,' Int. J. Forecasting, vol. 28, no. 1, pp. 57-66, 2012.",
        "C. Zhang et al., 'Forecasting realized volatility with spillover effects: Perspectives from graph neural networks,' Int. J. Forecasting, vol. 41, no. 1, pp. 377-397, 2025.",
        "R. Cheng and Q. Li, 'Modeling the momentum spillover effect for stock prediction via attribute-driven graph attention networks,' in Proc. AAAI Conf. Artif. Intell., vol. 35, no. 1, 2021.",
        "Y. Benjamini and Y. Hochberg, 'Controlling the false discovery rate: A practical and powerful approach to multiple testing,' J. Roy. Stat. Soc. B, vol. 57, no. 1, pp. 289-300, 1995.",
        "S. Hochreiter and J. Schmidhuber, 'Long short-term memory,' Neural Comput., vol. 9, no. 8, pp. 1735-1780, 1997.",
        "D. P. Kingma and J. Ba, 'Adam: A method for stochastic optimization,' in Proc. ICLR, 2015.",
        "F. X. Diebold and R. S. Mariano, 'Comparing predictive accuracy,' J. Bus. Econ. Stat., vol. 13, no. 3, pp. 253-263, 1995.",
        "D. Harvey, S. Leybourne, and P. Newbold, 'Testing the equality of prediction mean squared errors,' Int. J. Forecasting, vol. 13, no. 2, pp. 281-291, 1997.",
        "Q. McNemar, 'Note on the sampling error of the difference between correlated proportions or percentages,' Psychometrika, vol. 12, no. 2, pp. 153-157, 1947.",
        "M. H. Pesaran and A. Timmermann, 'A simple nonparametric test of predictive performance,' J. Bus. Econ. Stat., vol. 10, no. 4, pp. 461-465, 1992.",
        "R. S. Tsay, Analysis of Financial Time Series, 2nd ed. Hoboken, NJ, USA: Wiley, 2005.",
        "N. Meinshausen and P. Buhlmann, 'Stability selection,' J. Roy. Stat. Soc. B, vol. 72, no. 4, pp. 417-473, 2010.",
        "R. Cont, 'Empirical properties of asset returns: Stylized facts and statistical issues,' Quant. Finance, vol. 1, no. 2, pp. 223-236, 2001.",
    ]
    for idx, ref in enumerate(references, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(f"[{idx}] {ref}"), 8.0)

    doc.core_properties.title = "Identifying Granger-Causal Return Spillovers in Multivariate Financial Data with Interpretable Machine Learning"
    doc.core_properties.subject = "KSE 2026 camera-ready revision"
    doc.core_properties.author = "Nguyen Lam Tung Bach; Ngo Minh Huong; Nguyen Van Tang"
    doc.core_properties.keywords = "Neural Granger causality; return spillovers; VN-Index"
    doc.save(DOCX_PATH)
    print(f"Saved {DOCX_PATH}")


if __name__ == "__main__":
    build_document()
