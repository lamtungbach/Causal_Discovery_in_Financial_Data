from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent


def shade(cell, fill: str) -> None:
    prop = cell._tc.get_or_add_tcPr()
    node = prop.find(qn("w:shd")) or OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    if node.getparent() is None:
        prop.append(node)


def border_table(table) -> None:
    prop = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{name}")
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    prop.append(borders)


def add_table(doc: Document, lines: list[str]) -> None:
    data = []
    for line in lines:
        values = [x.strip() for x in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-+:?", x.replace(" ", "")) for x in values):
            continue
        data.append(values)
    table = doc.add_table(rows=len(data), cols=max(map(len, data)))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    border_table(table)
    for r, values in enumerate(data):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.size = Pt(8.3)
            if r == 0:
                shade(cell, "1F4E78")
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            elif r % 2 == 0:
                shade(cell, "EAF2F8")


def add_runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.*?\*\*)", text):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith("**") else part)
        run.bold = part.startswith("**")


def page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, value in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        run._r.append(node)


def build() -> Path:
    markdown = (HERE / "journal_manuscript.md").read_text(encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.82)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    lines = markdown.splitlines()
    i = 0
    title_done = False
    centered_lines = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i]); i += 1
            add_table(doc, table_lines)
            continue
        image = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line)
        if image:
            caption, relative = image.groups()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str((HERE / relative).resolve()), width=Inches(6.25))
            cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cp.add_run(caption); cr.italic = True; cr.font.size = Pt(9)
            i += 1; continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            level, text = len(heading.group(1)), heading.group(2)
            if level == 1 and not title_done:
                p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, text)
                for run in p.runs:
                    run.font.name = "Times New Roman"; run.font.size = Pt(16); run.font.color.rgb = RGBColor(0, 0, 0)
                title_done = True
            else:
                doc.add_heading(text, level=level)
            i += 1; continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, line[2:]); i += 1; continue
        p = doc.add_paragraph()
        if title_done and centered_lines < 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; centered_lines += 1
        add_runs(p, line)
        i += 1
    page_number(doc.sections[0].footer.paragraphs[0])
    output = HERE / "journal_manuscript.docx"
    doc.save(output)
    return output


if __name__ == "__main__":
    print(f"PASS: {build()}")
